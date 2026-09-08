#!/usr/bin/env python3
"""Builds tread_inspections_template.xlsx.

Formula constraint: this workbook is meant to be uploaded to Google Sheets, so it
uses only Excel-2007-era functions. MAXIFS/MINIFS are avoided because openpyxl must
write them as _xlfn.MAXIFS, which Google Sheets does not understand. SUMPRODUCT
gives the same results and is understood by Excel, LibreOffice and Google Sheets.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment

F = "Arial"
wb = Workbook()
hdr_font = Font(name=F, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor="1F2937")
body = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
title = Font(name=F, size=14, bold=True)
blue = Font(name=F, size=10, color="0000FF")          # input cells
green = Font(name=F, size=10, color="008000")         # link to another sheet
italic = Font(name=F, size=9, italic=True, color="555555")
yellow = PatternFill("solid", fgColor="FFFF00")
LAST = 2000   # data rows covered by the summary formulas

# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
content = [
    ("Scotia Tread Scanner — inspection workbook", title),
    ("", body),
    ("Getting this into Google Sheets", bold),
    ("1. Upload this file to Google Drive, open it, then File > Save as Google Sheets.", body),
    ("2. Delete the yellow EXAMPLE row on the Inspections tab once your first real data arrives.", body),
    ("", body),
    ("Adding an inspection", bold),
    ("1. In the app, finish the inspection and choose Share / download CSV.", body),
    ("2. In Google Sheets, click the Inspections tab, then File > Import > Upload.", body),
    ("3. Choose 'Append to current sheet' and 'Comma'. Leave 'Convert text to numbers' ticked.", body),
    ("The CSV columns are written in the same order as this tab, so append lines up every time.", body),
    ("", body),
    ("Fleet Summary", bold),
    ("Type a unit number into each yellow cell in column A. Every other cell on that tab is a formula.", body),
    ("The counts describe that unit's most recent inspection only, found by the highest inspection_id.", body),
    ("", body),
    ("What the columns mean", bold),
    ("depth_inner/centre/outer_32nds — the three groove readings in 32nds of an inch.", body),
    ("depth_min_32nds — the lowest of the three. This is what the status is based on.", body),
    ("status — OK, WATCH or REPLACE, decided by the app using the Thresholds tab values.", body),
    ("method — scan (LiDAR), gauge or manual. Keep this column: it is what makes a reading defensible.", body),
    ("position — TMC code. L or R for side, F/number/R for axle, O or I for outer or inner dual.", body),
    ("  Steer tires are LF and RF. Inner duals end in I and are always read with a gauge.", body),
    ("", body),
    ("Sources", bold),
    ("Tread minimums: Canada NSC Standard 11 and US FMCSA 49 CFR 393.75 — 4/32\" steer, 2/32\" all other positions.", body),
    ("Column layout: TREAD_SCANNER_PLAN.md in the ofakyle11/scotia repository.", body),
    ("Formulas cover Inspections rows 2 to %d. If you exceed that, widen the ranges on Fleet Summary." % LAST, body),
]
for i, (text, font) in enumerate(content, 1):
    c = ws.cell(row=i, column=1, value=text)
    c.font = font
    c.alignment = Alignment(wrap_text=False, vertical="top")
ws.column_dimensions["A"].width = 110
ws.sheet_view.showGridLines = False

# ---------------------------------------------------------------- Thresholds
th = wb.create_sheet("Thresholds")
th["A1"] = "Threshold settings"; th["A1"].font = title
rows = [
    ("Setting", "Value (32nds)", "Note"),
    ("Steer minimum", 4, "Legal minimum for steer axle (LF, RF)"),
    ("Drive / trailer minimum", 2, "Legal minimum for every other position"),
    ("Watch band above minimum", 2, "WATCH = within this many 32nds above the minimum"),
]
for r, row in enumerate(rows, 3):
    for c, v in enumerate(row, 1):
        cell = th.cell(row=r, column=c, value=v)
        if r == 3:
            cell.font = hdr_font; cell.fill = hdr_fill
        else:
            cell.font = blue if c == 2 else body
            if c == 2:
                cell.fill = yellow
th["B4"].comment = Comment("Source: Canada NSC Standard 11 / US FMCSA 49 CFR 393.75.", "Tread Scanner")
th["A8"] = "Yellow cells are inputs. The app applies its own copy of these thresholds when it writes the status column;"
th["A9"] = "these values are here so the workbook documents what the status column means."
for r in (8, 9):
    th.cell(row=r, column=1).font = italic
for col, w in zip("ABC", (28, 16, 58)):
    th.column_dimensions[col].width = w
th.sheet_view.showGridLines = False

# ---------------------------------------------------------------- Inspections
ins = wb.create_sheet("Inspections")
header = ["inspection_id", "date", "technician", "customer", "unit_number", "plate", "vin", "odometer",
          "axle_config", "position", "brand", "model", "size", "dot_code",
          "depth_inner_32nds", "depth_centre_32nds", "depth_outer_32nds", "depth_min_32nds", "depth_min_mm",
          "pressure_psi", "status", "method", "photo_url", "notes", "scan_confidence_32nds"]
for c, name in enumerate(header, 1):
    cell = ins.cell(row=1, column=c, value=name)
    cell.font = hdr_font; cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
example = ["20260908-1432-AB12", "2026-09-08T14:32:00", "M. Francis", "Acme Freight", "42", "ABC 123",
           "1FUJGLDR5CSBB1234", 512000, "TRACTOR_3A", "LF", "Michelin", "XZE2", "11R22.5", "DOT B9 4K 2325",
           6, 5, 7, 5, 3.97, 105, "WATCH", "gauge", "", "EXAMPLE ROW — delete once real data is imported", ""]
for c, v in enumerate(example, 1):
    cell = ins.cell(row=2, column=c, value=v)
    cell.font = body; cell.fill = yellow
ins.cell(row=2, column=1).comment = Comment(
    "Example row showing the expected format. Delete it after your first real import.\n"
    "inspection_id is yyyymmdd-hhmm-XXXX and sorts chronologically, which Fleet Summary relies on.",
    "Tread Scanner")
widths = [19, 19, 13, 17, 11, 10, 19, 10, 12, 9, 11, 10, 11, 15, 9, 9, 9, 9, 9, 8, 10, 8, 13, 34, 10]
for i, w in enumerate(widths, 1):
    ins.column_dimensions[get_column_letter(i)].width = w
ins.row_dimensions[1].height = 30
ins.freeze_panes = "A2"
ins.auto_filter.ref = "A1:Y%d" % LAST
for op, colr in (("equal", "FCA5A5"), ("equal", "FDE68A"), ("equal", "BBF7D0")):
    pass
ins.conditional_formatting.add("U2:U%d" % LAST, CellIsRule(operator="equal", formula=['"REPLACE"'],
    fill=PatternFill("solid", fgColor="FCA5A5"), font=Font(name=F, size=10, bold=True, color="7F1D1D")))
ins.conditional_formatting.add("U2:U%d" % LAST, CellIsRule(operator="equal", formula=['"WATCH"'],
    fill=PatternFill("solid", fgColor="FDE68A"), font=Font(name=F, size=10, color="78350F")))
ins.conditional_formatting.add("U2:U%d" % LAST, CellIsRule(operator="equal", formula=['"OK"'],
    fill=PatternFill("solid", fgColor="BBF7D0"), font=Font(name=F, size=10, color="14532D")))

# ---------------------------------------------------------------- Fleet Summary
fs = wb.create_sheet("Fleet Summary")
fs["A1"] = "Fleet summary — most recent inspection per unit"; fs["A1"].font = title
fs["A2"] = "Type unit numbers into the yellow cells. Everything else is calculated."; fs["A2"].font = italic
fh = ["Unit", "Customer", "Last inspection", "Odometer", "Latest ref", "Tires",
      "Replace", "Watch", "OK", "Lowest tread", "Lowest position", "Action"]
HROW = 4
for c, name in enumerate(fh, 1):
    cell = fs.cell(row=HROW, column=c, value=name)
    cell.font = hdr_font; cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal="center", wrap_text=True)

I = "Inspections!"
ID  = f"{I}$A$2:$A${LAST}"
DT  = f"{I}$B$2:$B${LAST}"
CU  = f"{I}$D$2:$D${LAST}"
UN  = f"{I}$E$2:$E${LAST}"
OD  = f"{I}$H$2:$H${LAST}"
POS = f"{I}$J$2:$J${LAST}"
MIN = f"{I}$R$2:$R${LAST}"
ST  = f"{I}$U$2:$U${LAST}"
# Sortable numeric key from inspection_id "yyyymmdd-hhmm-XXXX" -> yyyymmddhhmm
KEY = f'IFERROR(VALUE(LEFT({ID},8)&MID({ID},10,4)),0)'

FIRST, ROWS = HROW + 1, 40
for i in range(ROWS):
    r = FIRST + i
    u = f"$A{r}"
    blank = f'{u}=""'
    a = fs.cell(row=r, column=1); a.fill = yellow; a.font = blue
    # E: inspection_id of that unit's latest inspection (highest key)
    # Latest inspection_id for this unit. SUMPRODUCT(MAX(cond*ROW())) needs no array entry and
    # is understood identically by Excel, LibreOffice and Google Sheets.
    fs.cell(row=r, column=5, value=(
        f'=IF({blank},"",IFERROR(INDEX({ID},SUMPRODUCT(MAX(({UN}={u})*'
        f'({KEY}=SUMPRODUCT(MAX(({UN}={u})*{KEY})))*ROW({ID})))-1),""))'))
    ref = f"$E{r}"
    ok = f'IF(OR({blank},{ref}=""),""'
    fs.cell(row=r, column=2, value=f'={ok},IFERROR(INDEX({CU},MATCH({ref},{ID},0)),""))')
    fs.cell(row=r, column=3, value=f'={ok},IFERROR(INDEX({DT},MATCH({ref},{ID},0)),""))')
    fs.cell(row=r, column=4, value=f'={ok},IFERROR(INDEX({OD},MATCH({ref},{ID},0)),""))')
    fs.cell(row=r, column=6, value=f'={ok},COUNTIFS({ID},{ref},{MIN},"<>"))')
    for col, status in ((7, "REPLACE"), (8, "WATCH"), (9, "OK")):
        fs.cell(row=r, column=col, value=f'={ok},COUNTIFS({ID},{ref},{ST},"{status}"))')
    # Lowest tread without MINIFS: max of (99 - depth) over matching rows, inverted.
    fs.cell(row=r, column=10, value=(
        f'={ok},IF($F{r}=0,"",99-SUMPRODUCT(MAX(({ID}={ref})*({MIN}<>"")*(99-{MIN})))))'))
    fs.cell(row=r, column=11, value=(
        f'=IF(OR({blank},$J{r}=""),"",IFERROR(INDEX({POS},'
        f'SUMPRODUCT(MAX(({ID}={ref})*({MIN}=$J{r})*ROW({POS})))-1),""))'))
    fs.cell(row=r, column=12, value=(
        f'={ok},IF($G{r}>0,"Replace "&$G{r}&" tire(s) now — unit is out of service",'
        f'IF($H{r}>0,"Book "&$H{r}&" tire(s) soon","OK")))'))
    for c in range(2, 13):
        fs.cell(row=r, column=c).font = green if c in (2, 3, 4, 5) else body
    fs.cell(row=r, column=10).font = bold
fs["A%d" % FIRST] = "42"   # matches the example row so the tab is not empty on open

for i, w in enumerate([10, 18, 20, 12, 20, 8, 10, 9, 7, 13, 15, 44], 1):
    fs.column_dimensions[get_column_letter(i)].width = w
fs.freeze_panes = "B%d" % FIRST
fs.conditional_formatting.add(f"G{FIRST}:G{FIRST+ROWS-1}", CellIsRule(operator="greaterThan", formula=["0"],
    fill=PatternFill("solid", fgColor="FCA5A5"), font=Font(name=F, size=10, bold=True, color="7F1D1D")))
fs.conditional_formatting.add(f"H{FIRST}:H{FIRST+ROWS-1}", CellIsRule(operator="greaterThan", formula=["0"],
    fill=PatternFill("solid", fgColor="FDE68A")))
note = FIRST + ROWS + 1
fs.cell(row=note, column=1, value="Counts describe the latest inspection for each unit, identified by the highest inspection_id.").font = italic
fs.cell(row=note + 1, column=1, value="If a unit is inspected twice in the same minute, only the later reference is used.").font = italic
fs.cell(row=note + 2, column=1, value="Lowest tread is the smallest depth_min_32nds on that inspection, in 32nds of an inch.").font = italic
fs.sheet_view.showGridLines = False

wb.move_sheet("Inspections", offset=-2)
wb.move_sheet("Fleet Summary", offset=-2)
wb.save("tread_inspections_template.xlsx")
print("saved; sheet order:", wb.sheetnames)
